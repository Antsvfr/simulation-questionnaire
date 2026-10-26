from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
IND=RGBColor(0x4F,0x46,0xE5); AMB=RGBColor(0xF5,0x9E,0x0B); DK=RGBColor(0x1F,0x29,0x37); W=RGBColor(255,255,255); LG=RGBColor(0xEE,0xF2,0xFF)
p=Presentation(); p.slide_width=Inches(13.333); p.slide_height=Inches(7.5)
B=p.slide_layouts[6]
def tb(s,x,y,w,h,text,size=18,color=DK,bold=False):
    t=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)).text_frame; t.word_wrap=True
    for i,l in enumerate(text.split("\n")):
        para=t.paragraphs[0] if i==0 else t.add_paragraph()
        r=para.add_run(); r.text=l; r.font.size=Pt(size); r.font.color.rgb=color; r.font.bold=bold; para.space_after=Pt(6)
def rect(s,x,y,w,h,c,shape=MSO_SHAPE.RECTANGLE):
    r=s.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h)); r.fill.solid(); r.fill.fore_color.rgb=c; r.line.fill.background(); return r
def slide(title,section=None):
    s=p.slides.add_slide(B); rect(s,0,0,13.333,1.2,IND); rect(s,0,1.2,13.333,0.08,AMB)
    tb(s,0.6,0.2,10,0.9,title,32,W,True)
    if section: tb(s,9.8,0.35,3.2,0.6,section,14,AMB,True)
    return s
def bullets(s,items,x=0.7,y=1.7,w=12,size=20): tb(s,x,y,w,5.5,"\n".join("• "+i for i in items),size)
def card(s,x,y,w,h,title,body):
    rect(s,x,y,w,h,LG,MSO_SHAPE.ROUNDED_RECTANGLE); tb(s,x+0.2,y+0.15,w-0.4,0.6,title,18,IND,True); tb(s,x+0.2,y+0.8,w-0.4,h-0.9,body,14)
def shot(s,x,y,w,h,label):
    r=rect(s,x,y,w,h,RGBColor(0xF3,0xF4,0xF6)); r.line.color.rgb=RGBColor(0x9C,0xA3,0xAF); r.line.dash_style=4
    tb(s,x+0.2,y+h/2-0.4,w-0.4,0.8,"📸 Capture d'écran à insérer :\n"+label,13,RGBColor(0x6B,0x72,0x80))
# 1
s=p.slides.add_slide(B); rect(s,0,0,13.333,7.5,IND); rect(s,0,5.2,13.333,0.12,AMB)
tb(s,0.8,1.8,12,1.5,"AlterMatch",66,W,True); tb(s,0.8,3.3,12,1,"Chaque alternance, un vrai match.",28,AMB)
tb(s,0.8,5.6,12,1.2,"Final Challenge — Start-up & Chatbot Zapier\n[Votre nom] · [Prénom] · [Prénom]",18,W)
# 2
s=slide("L'équipe","START-UP")
for i,(n,r) in enumerate([("[Votre nom]","CEO & Stratégie"),("[Prénom]","Produit & Chatbot"),("[Prénom]","Growth & Partenariats écoles")]):
    x=0.9+i*4.1; rect(s,x+1,1.9,1.6,1.6,LG,MSO_SHAPE.OVAL); tb(s,x,3.8,3.6,0.6,n,22,DK,True); tb(s,x,4.4,3.6,0.6,r,16,IND)
# 3
s=slide("Le problème","START-UP")
card(s,0.7,1.7,5.8,4.8,"🎓 Étudiant·e·s","Des dizaines de candidatures sans réponse\nRecherche qui dure des mois\nRythme école mal compris des entreprises\nStress et risque de perdre sa place en école")
card(s,6.85,1.7,5.8,4.8,"🏢 PME (10–250 salariés)","Pas de service RH pour trier\nAides et démarches (OPCO, CERFA) floues\nRecrutements tardifs, profils mal adaptés\nRuptures de contrat pendant l'essai")
# 4
s=slide("Mission & proposition de valeur","START-UP")
tb(s,0.7,1.6,12,0.8,"Mission : faire de chaque alternance un vrai match, pour l'étudiant comme pour l'entreprise.",20,IND,True)
for i,(t,b) in enumerate([("Matching 3D","Compétences + rythme d'alternance + culture d'entreprise"),("Shortlist en 72 h","3 profils pré-qualifiés, entretiens organisés"),("Zéro paperasse","CERFA, OPCO et aides gérés par AlterMatch"),("Assistant IA 24/7","Alfie qualifie le besoin en 3 minutes")]):
    card(s,0.7+i*3.1,2.7,2.9,3.6,t,b)
# 5
s=slide("Clientèle cible","START-UP")
card(s,0.7,1.7,5.8,4.8,"Clients payeurs : PME","Dirigeant·e·s / managers, France\nProblèmes : temps, administratif, besoin de profils opérationnels\nComportement : cherchent sur Google le soir, veulent une réponse immédiate\nPersona : Claire, 45 ans, dirige une agence de 40 salariés")
card(s,6.85,1.7,5.8,4.8,"Utilisateurs : étudiant·e·s","19–25 ans, Bachelor / Master école de commerce\nProblèmes : visibilité, stress, calendrier\nComportement : mobile first, attendent des réponses 24/7\nPersona : Yanis, 21 ans, cherche une alternance marketing")
# 6
s=slide("Modèle de revenus","START-UP")
for i,(t,pr,b) in enumerate([("Étudiant·e·s","0 €","Gratuit, toujours"),("Offre Succès","990 € HT","par alternant signé\nrien si pas de recrutement"),("Offre Pro","149 € HT/mois","recrutements illimités\ntableau de bord\naccompagnement admin"),("Écoles","Licence annuelle","suivi des placements")]):
    x=0.7+i*3.1; rect(s,x,1.8,2.9,4.6,LG,MSO_SHAPE.ROUNDED_RECTANGLE); tb(s,x+0.2,2.0,2.5,0.6,t,18,DK,True); tb(s,x+0.2,2.7,2.5,0.9,pr,26,IND,True); tb(s,x+0.2,3.8,2.5,2.4,b,14)
tb(s,0.7,6.6,12,0.6,"Garantie : rupture pendant la période d'essai = nouveau profil offert.",16,AMB,True)
# 7
s=slide("Stratégie chatbot : pourquoi Alfie ?","CHATBOT")
tb(s,0.7,1.6,12,0.8,"Objectif : qualification de prospects + prise de rendez-vous (FAQ étudiant·e·s en secondaire)",20,IND,True)
bullets(s,["Les dirigeant·e·s cherchent le soir : réponse instantanée 24/7","Chaque visiteur entreprise devient un lead qualifié (taille, poste, date, ville)","L'équipe commerciale ne traite que les leads chauds","Les questions répétitives (tarifs, aides) sont automatisées","KPI visés : conversion visiteur→lead x2 · 70 % des FAQ sans humain · RDV < 24 h"],y=2.5)
# 8
s=slide("Parcours utilisateur","CHATBOT")
steps=["Accueil","Entreprise\nou étudiant ?","6 questions\n(1 à la fois)","Résumé +\nvalidation","Coordonnées\n+ appel","Zap :\nfiltre + paths"]
for i,t in enumerate(steps):
    x=0.5+i*2.12; r=rect(s,x,2.8,1.85,1.6,IND if i%2==0 else AMB,MSO_SHAPE.ROUNDED_RECTANGLE); tb(s,x+0.1,3.1,1.65,1.2,t,15,W,True)
    if i<5: tb(s,x+1.85,3.2,0.3,0.6,"→",24,DK,True)
tb(s,0.7,5.0,12,1.5,"Questions : entreprise & secteur → taille → poste → niveau & rythme → date de démarrage → ville / télétravail\nBranche étudiant : service gratuit + lien d'inscription altermatch.fr/etudiant",15)
# 9
s=slide("Mise en œuvre 1 : accueil & directives","MISE EN ŒUVRE")
tb(s,0.7,1.6,6,4.8,"Message d'accueil :\n« 👋 Bonjour, je suis Alfie, l'assistant d'AlterMatch ! Je vous aide à trouver l'alternant·e idéal·e en moins de 72 h — ou à décrocher votre alternance. Vous êtes plutôt entreprise 🏢 ou étudiant·e 🎓 ? »\n\nDirectives : rôle · objectifs · étapes · règles · audience · contexte · ton\nRègles clés : 1 question à la fois, max 5 lignes, jamais de prix/aide inventés, escalade humaine immédiate, pas de données sensibles",15)
shot(s,7.0,1.6,5.7,5.2,"Éditeur Zapier — Directive & Greeting")
# 10
s=slide("Mise en œuvre 2 : connaissance, thème, modèle","MISE EN ŒUVRE")
tb(s,0.7,1.6,6,5,"Knowledge source : FAQ AlterMatch (fonctionnement, tarifs, aides, rythmes, contact) — réponses limitées à ce contenu\n\nThème : indigo #4F46E5 + ambre #F59E0B, avatar « AM »\n\nModèle d'IA : modèle par défaut Zapier, créativité basse pour la fiabilité\n\nBoutons de démarrage : « Je cherche un alternant » · « Combien ça coûte ? » · « Quelles aides ? » · « Je suis étudiant·e »",15)
shot(s,7.0,1.6,2.75,2.5,"Knowledge"); shot(s,9.95,1.6,2.75,2.5,"Theme"); shot(s,7.0,4.3,5.7,2.5,"Modèle d'IA + Lead form")
# 11
s=slide("Workflow Zapier : filtres & conditions","MISE EN ŒUVRE")
for i,(t,c) in enumerate([("Trigger\nNew Lead (Chatbot)",IND),("Filter\n≥ 10 salariés\nET date renseignée",AMB),("Path A — chaud\ndémarrage < 2 mois\n→ Gmail + Slack",IND),("Path B — tiède\n→ Google Sheets\n+ email nurturing",IND)]):
    x=0.6+i*3.15; rect(s,x,1.8,2.8,1.9,c,MSO_SHAPE.ROUNDED_RECTANGLE); tb(s,x+0.15,1.95,2.5,1.7,t,14,W,True)
shot(s,0.6,4.1,12.1,2.9,"Éditeur du Zap (trigger, filter, paths) + historique Zap Runs")
# 12
s=slide("Démo : Alfie en action","MISE EN ŒUVRE")
for i,l in enumerate(["Qualification PME 40 salariés","Question tarifs","Question aides (pas d'invention)","« Je veux parler à un humain »"]):
    shot(s,0.6+i*3.1,1.6,2.9,4.4,l)
tb(s,0.6,6.3,12,0.7,"🔗 Lien du chatbot : [coller le lien Zapier ici]",20,IND,True)
# 13
s=slide("Notre avantage concurrentiel","PITCH")
card(s,0.7,1.7,5.8,4.8,"Jobboards classiques","Annonce passive, réponse en jours\nTri manuel des CV\nAucune aide administrative\nCoût par lead élevé")
card(s,6.85,1.7,5.8,4.8,"AlterMatch + Alfie","Réponse instantanée 24/7\nBesoin qualifié automatiquement\nAides & démarches expliquées\nHumain concentré sur les leads chauds → CAC réduit")
# 14
s=slide("Prochaines étapes","PITCH")
bullets(s,["Prise de RDV directe (Calendly) dans le chat","Chatbot sur WhatsApp & LinkedIn","Scoring automatique des leads","Version étudiant : coaching CV et entretiens"],y=1.8,size=22)
tb(s,0.7,5.5,12,1,"Merci ! Questions ?",40,IND,True)
p.save("FinalChallenge_VotreNom.pptx")
